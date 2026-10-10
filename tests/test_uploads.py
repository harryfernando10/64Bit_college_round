import io
import zipfile

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from backend.main import MAX_UPLOAD_BYTES, app


client = TestClient(app)


def make_pdf(text: str) -> bytes:
    # Minimal single-page PDF with an uncompressed text stream.
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(pdf)


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("The source confirms 48 plants.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Trial duration"
    table.cell(0, 1).text = "Four weeks"
    content = io.BytesIO()
    document.save(content)
    return content.getvalue()


def upload(filename: str, content: bytes, media_type: str = "application/octet-stream"):
    return client.post("/api/extract-document", files={"file": (filename, content, media_type)})


def test_upload_extracts_utf8_text_and_preserves_filename():
    response = upload("field-notes.txt", b"\xef\xbb\xbfGarden trial ended after four weeks.", "text/plain")
    assert response.status_code == 200
    assert response.json()["filename"] == "field-notes.txt"
    assert response.json()["file_type"] == "TXT"
    assert response.json()["text"] == "Garden trial ended after four weeks."


def test_upload_extracts_text_pdf():
    response = upload("research.pdf", make_pdf("The source confirms 48 plants."), "application/pdf")
    assert response.status_code == 200
    assert "The source confirms 48 plants." in response.json()["text"]


def test_upload_extracts_docx_paragraphs_and_table_cells():
    response = upload("research.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.status_code == 200
    assert response.json()["text"] == "The source confirms 48 plants.\nTrial duration\tFour weeks"


def test_extracted_text_flows_into_existing_audit_and_combines_with_manual_sources():
    extracted = upload("research.txt", b"Researchers counted 48 plants in the trial.", "text/plain").json()
    sources = [
        {"source_id": "upload-1", "title": extracted["filename"], "text": extracted["text"]},
        {"source_id": "manual-1", "title": "Manual note", "text": "The field was near the river."},
    ]
    response = client.post("/api/audit", json={"answer": "Researchers counted 84 plants in the trial.", "sources": sources})
    assert response.status_code == 200
    data = response.json()
    assert data["execution_mode"] == "BASELINE"
    assert data["sources"] == sources
    assert data["claims"][0]["verdict"] == "CONTRADICTED"
    reference = data["claims"][0]["evidence"][0]
    assert extracted["text"][reference["start_offset"]:reference["end_offset"]] == reference["quote"]


def test_empty_unsupported_and_malformed_files_return_readable_errors():
    missing = client.post("/api/extract-document")
    assert missing.status_code == 422
    assert upload("empty.txt", b"").status_code == 422
    unsupported = upload("notes.rtf", b"Some text")
    assert unsupported.status_code == 415 and "Unsupported" in unsupported.json()["detail"]
    bad_pdf = upload("broken.pdf", b"not a pdf")
    assert bad_pdf.status_code == 422 and "corrupt" in bad_pdf.json()["detail"]
    bad_docx = upload("broken.docx", b"PKnot a document")
    assert bad_docx.status_code == 422 and "corrupt" in bad_docx.json()["detail"]
    mislabeled = upload("mislabelled.txt", make_pdf("not plain text"))
    assert mislabeled.status_code == 422 and "do not match" in mislabeled.json()["detail"]
    binary_text = upload("binary.txt", b"hello\x00world")
    assert binary_text.status_code == 422 and "control characters" in binary_text.json()["detail"]
    oversized_text = upload("too-much-text.txt", b"x" * 30001)
    assert oversized_text.status_code == 413 and "30,000-character" in oversized_text.json()["detail"]
    empty_extraction = upload("scanned.pdf", make_pdf(""), "application/pdf")
    assert empty_extraction.status_code == 422 and "OCR is not supported" in empty_extraction.json()["detail"]


def test_encrypted_pdf_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("password")
    content = io.BytesIO()
    writer.write(content)
    response = upload("locked.pdf", content.getvalue(), "application/pdf")
    assert response.status_code == 422 and "Encrypted" in response.json()["detail"]


def test_file_size_limit_is_enforced_and_filename_is_sanitized():
    oversized = upload("too-large.txt", b"x" * (MAX_UPLOAD_BYTES + 1), "text/plain")
    assert oversized.status_code == 413
    response = upload(r"C:\private\report.txt", b"Short report.", "text/plain")
    assert response.status_code == 200
    assert response.json()["filename"] == "report.txt"


def test_upload_source_counts_share_the_existing_eight_source_limit():
    uploaded = upload("notes.txt", b"Source note.", "text/plain").json()
    sources = [{"source_id": "upload-1", "title": uploaded["filename"], "text": uploaded["text"]}]
    sources.extend({"source_id": f"manual-{index}", "title": f"Manual {index}", "text": f"Manual source {index}."} for index in range(2, 9))
    answer = "A custom claim."
    assert client.post("/api/audit", json={"answer": answer, "sources": sources}).status_code == 200
    assert client.post("/api/audit", json={"answer": answer, "sources": [*sources, {"source_id": "manual-9", "title": "Manual 9", "text": "One too many."}]}).status_code == 422


def test_failed_upload_does_not_change_a_valid_upload_result():
    valid = upload("valid.txt", b"Retained evidence.", "text/plain").json()
    invalid = upload("bad.pdf", b"invalid")
    assert invalid.status_code == 422
    audit = client.post("/api/audit", json={"answer": "Retained evidence.", "sources": [{"source_id": "kept", "title": valid["filename"], "text": valid["text"]}]})
    assert audit.status_code == 200
    assert audit.json()["claims"][0]["verdict"] == "SUPPORTED"


def test_docx_corrupt_zip_is_rejected():
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
    response = upload("incomplete.docx", content.getvalue())
    assert response.status_code == 422
